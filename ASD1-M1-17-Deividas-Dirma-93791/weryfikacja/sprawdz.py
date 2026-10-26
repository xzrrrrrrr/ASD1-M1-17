"""Kontrola paczki bez zmiany kodu. Python 3 + pymupdf + python-pptx.
Dodatkowo --przegladarka wymaga selenium, Firefoksa i geckodriver w PATH.
Przykład: python3 weryfikacja/sprawdz.py --przegladarka --raport /tmp/wyniki-asd.json
Odbiorca aplikacji nie potrzebuje tych narzędzi: wystarczy index.html lub tests.html.
"""
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urlsplit,unquote
from zipfile import ZipFile
import argparse,json,re,shutil,tempfile,hashlib,os,posixpath

# Domyślnie sprawdzamy folder nad katalogiem weryfikacja; można wskazać inną kopię.
ap=argparse.ArgumentParser();ap.add_argument('--projekt',type=Path,default=Path(__file__).resolve().parents[1]);ap.add_argument('--przegladarka',action='store_true');ap.add_argument('--raport',type=Path);args=ap.parse_args();root=args.projekt.resolve()
class Refs(HTMLParser):
    def __init__(self):super().__init__();self.refs=[]
    def handle_starttag(self,tag,attrs):
        for k,v in attrs:
            if k in ['src','href'] and v:self.refs.append(v)
local=[]
for p in root.rglob('*.html'):
    parser=Refs();parser.feed(p.read_text())
    for ref in parser.refs:
        u=urlsplit(ref)
        if u.scheme or u.netloc or not u.path:continue
        target=(p.parent/unquote(u.path)).resolve();assert target.exists(),(p,ref);local.append(str(target.relative_to(root)))
for ref in re.findall(r'url\([\s\'"]*([^\)\'\"]+)',(root/'style.css').read_text()):
    if not urlsplit(ref).scheme:assert (root/ref).exists()
for ref in re.findall(r'!\[[^\]]*\]\(([^)]+)\)',(root/'DOKUMENTACJA.md').read_text()):assert (root/ref).is_file()
for ref in re.findall(r'`([^`]+)`',(root/'README.md').read_text()):
    if re.match(r'^[\w/-]+\.(html|css|js|pptx|pdf|docx|json|md)$',ref) or ref.endswith('/'):assert (root/ref).exists(),ref
# Liczymy rzeczywiste strony i slajdy, a nie nagłówki w źródle.
import pymupdf as fitz
from pptx import Presentation
prs=Presentation(root/'prezentacja.pptx');assert len(prs.slides)==10
counts={}
for name in ['DOKUMENTACJA.pdf','prezentacja.pdf']:
    pdf=fitz.open(root/name);assert len(pdf)==10;counts[name]=len(pdf)
    for pg in pdf:
        assert pg.get_text().strip()
        for b in pg.get_text('dict')['blocks']:
            if b['type']!=0:continue
            for line in b['lines']:
                for span in line['spans']:
                    x0,y0,x1,y1=span['bbox'];assert 0<=x0<x1<=pg.rect.width+1 and 0<=y0<y1<=pg.rect.height+1
for name in ['prezentacja.pptx','DOKUMENTACJA.docx']:
    with ZipFile(root/name) as z:
        assert z.testzip() is None
        for rel in [n for n in z.namelist() if n.endswith('.rels')]:
            import xml.etree.ElementTree as ET
            for x in ET.fromstring(z.read(rel)):
                if x.attrib.get('TargetMode')=='External':continue
                target=x.attrib['Target'];base=posixpath.dirname(posixpath.dirname(rel))
                dest=posixpath.normpath(posixpath.join(base,target)).lstrip('/')
                assert dest in z.namelist(),(name,rel,target,dest)
readme=(root/'README.md').read_text();doc=(root/'DOKUMENTACJA.md').read_text();slide_text='\n'.join(s.text for sl in prs.slides for s in sl.shapes if s.has_text_frame);notes='\n'.join(sl.notes_slide.notes_text_frame.text for sl in prs.slides)
quote='Przy realizacji projektu korzystałem z agenta OpenClaw, do którego podłączony był Codex Sol 6.1 z ustawieniem rozumowania medium'
assert quote in readme and quote in doc and quote in notes
for t in [readme,doc,slide_text]:
    assert 'Deividas Dirma' in t and '93791' in t
    for bad in ['GPT-6','do uzupełnienia','TODO','placeholder','Lorem ipsum']:assert bad not in t,bad
checks=[]
def check(name,expected,actual):
    r={'nazwa':name,'oczekiwane':expected,'rzeczywiste':actual,'pass':expected==actual};checks.append(r);assert r['pass'],r
result={'wykonawca':'agent lub operator narzędzia','localReferences':local,'slajdy':len(prs.slides),'pages':counts,'notes':len(prs.slides),'kontroleUI':checks}
if args.przegladarka:
    from selenium import webdriver
    from selenium.webdriver.firefox.service import Service
    from selenium.webdriver.support.ui import Select
    # Snap Firefox może nie widzieć katalogu ukrytego. Kopia jest porównywana bajtowo.
    browser_root=root
    if any(part.startswith('.') for part in root.parts if part not in ['.','..']):
        browser_root=Path(tempfile.mkdtemp(prefix='asd-kontrola-',dir=Path.home()))
        for f in ['index.html','tests.html','style.css','algorithms.js','script.js','animation.js','education.js','comparison.js','tests.js']:
            shutil.copy2(root/f,browser_root/f);assert (root/f).read_bytes()==(browser_root/f).read_bytes()
    opts=webdriver.FirefoxOptions();opts.add_argument('-headless')
    for k,v in {'network.proxy.type':1,'network.proxy.http':'127.0.0.1','network.proxy.http_port':9,'network.proxy.ssl':'127.0.0.1','network.proxy.ssl_port':9,'network.proxy.no_proxies_on':''}.items():opts.set_preference(k,v)
    driver_path=shutil.which('geckodriver');assert driver_path,'Brak geckodriver w PATH'
    d=webdriver.Firefox(options=opts,service=Service(driver_path))
    def js(code):return d.execute_script('return window.eval(arguments[0]);',code)
    try:
        d.set_window_size(1440,1150);d.get((browser_root/'tests.html').as_uri())
        assert d.find_element('id','status-testow').text=='PASS 46, FAIL 0'
        tests=js('uruchomTestyAlgorytmow()');assert tests['pass']==46 and tests['fail']==0;result['testy']=tests
        d.get((browser_root/'index.html').as_uri())
        check('Przykład karty',[1,6,2],js('[Number(document.getElementById("liczba-konfliktow").textContent),Number(document.getElementById("dlugosc-pokrycia").textContent),Number(document.getElementById("wartosc-maksimum").textContent)]'))
        for name,spans,expected in [('Karta',[[1,4],[2,5],[5,7]],{'maksimum':2,'odcinki':[{'start':2,'end':4}]}),('Rozłączne',[[10,30],[20,40],[60,80],[70,90]],{'maksimum':2,'odcinki':[{'start':20,'end':30},{'start':70,'end':80}]}),('Styk',[[1,4],[4,7]],{'maksimum':1,'odcinki':[{'start':1,'end':7}]}),('Puste',[],{'maksimum':0,'odcinki':[]})]:
            data=[{'id':chr(65+i),'start':a,'end':b} for i,(a,b) in enumerate(spans)]
            js('przedzialy='+json.dumps(data)+';odswiezWidok();')
            actual=js('({maksimum:Number(document.getElementById("wartosc-maksimum").textContent),odcinki:Array.from(document.querySelectorAll("#os-maksimum .maksimum-koncowe")).map(p=>({start:parseFloat(p.style.left),end:Math.round((parseFloat(p.style.left)+parseFloat(p.style.width))*100)/100}))})')
            check('Końcowa oś: '+name,expected,actual)
        js('document.getElementById("scenariusz").value="1";wczytajScenariusz();')
        for method in ['konflikty','laczenie','zdarzenia','maksimum','jednostki']:
            Select(d.find_element('id','algorytm')).select_by_value(method);d.find_element('id','krok').click();check('Krok: '+method,True,js('numerKroku===0&&!!document.querySelector("#pseudokod .aktywna-linia")&&document.getElementById("opis-kroku").textContent===krokiAnimacji[0].stan.opis'))
        js('document.getElementById("edycja-danych").open=true;')
        for ident,val in [('start','12'),('end','18')]:
            el=d.find_element('id',ident);el.clear();el.send_keys(val)
        d.find_element('id','zapisz-przedzial').click();check('Dodanie',4,js('przedzialy.length'))
        for ident,val in [('start','4'),('end','4')]:
            el=d.find_element('id',ident);el.clear();el.send_keys(val)
        d.find_element('id','zapisz-przedzial').click();check('Walidacja bez zmiany danych',True,js('przedzialy.length===4&&document.getElementById("komunikat").classList.contains("blad")'))
        d.find_element('css selector','#lista-przedzialow tr:last-child button:first-child').click();el=d.find_element('id','end');el.clear();el.send_keys('19');d.find_element('id','zapisz-przedzial').click();check('Edycja',19,js('przedzialy[3].end'))
        d.find_element('css selector','#lista-przedzialow tr:last-child .usun').click();check('Usunięcie',3,js('przedzialy.length'))
        js('document.getElementById("scenariusz").value="3";wczytajScenariusz();wykonajKrok();resetujAnimacje();');check('Reset zachowuje maksimum',[2,-1],js('[document.querySelectorAll("#os-maksimum .maksimum-koncowe").length,numerKroku]'))
        d.find_element('id','start-pauza').click();running=js('zegarAnimacji!==null');d.find_element('id','start-pauza').click();check('Start/Pauza',True,running and js('zegarAnimacji===null'))
        result['porownanie']=js('[5,15,30].map(porownajMetody)');assert all(x['zgodne'] for x in result['porownanie'])
        result['browser']=d.capabilities['browserVersion'];result['tryb']='file://; HTTP/HTTPS przez zamknięty port proxy 127.0.0.1:9';result['browserRoot']=str(browser_root)
    finally:d.quit()
if args.raport:args.raport.write_text(json.dumps(result,ensure_ascii=False,indent=2))
print('PPTX 10 slajdów; PDF-y po 10 stron; lokalne odwołania poprawne.')
if args.przegladarka:print('Testy: PASS 46, FAIL 0; oddzielnie UI '+str(sum(x['pass'] for x in checks))+'/'+str(len(checks))+'; porównanie 5/15/30 zgodne.')
